"""个人文档库：自有文档（PDF/Word/TXT/MD）→ 检索 → 自由问答。

与评测线（fixture + 选项判定）的关系：共用检索层（BM25F + 分块）与模型
客户端（token 台账），答题头换成**自由作答**——个人场景没有选项可判，
答案是自由文本，但纪律不变：

- 只许依据文档作答，文档没有就明说「未提及」（不编造）；
- 答案必须带【文档·Section】出处引用，引用经在场匹配校验（M4 归因规则）；
- token 台账覆盖全部调用。

用户数据一律放 ``findata_userdata/``（gitignore）：documents/ 存原件，
converted/ 存转换后的 markdown（`## ` 分节锚点供检索层切块）。
"""

from __future__ import annotations

import re
from pathlib import Path

import pymupdf

from findata.finqa.attribution import extract_citations, match_citations
from findata.finqa.baseline import MAX_CHUNK_CHARS
from findata.finqa.embedding import Embedder, EmbeddingUnavailable
from findata.finqa.hybrid import HybridIndex
from findata.finqa.qwen import QwenClient
from findata.finqa.retrieval import BM25Index, chunk_docs
from findata.finqa.tables import render_markdown_table

SAFE_NAME = re.compile(r"[^A-Za-z0-9_\-\u4e00-\u9fff.]+")
_WS = re.compile(r"[ \t]+")
_BLANK = re.compile(r"\n\s*\n+")
_ANSWER = re.compile(r"ANSWER\s*[:：]\s*(.+)")


def safe_stem(filename: str) -> str:
    stem = Path(filename).stem
    cleaned = SAFE_NAME.sub("_", stem).strip("._") or "document"
    return cleaned[:80]


# ── 转换 ──


def _page_blocks_outside_tables(page: pymupdf.Page) -> list[str]:
    """页面文本块，剔除落在表格 bbox 内的块（表格文本由专用渲染产出，
    不剔除会整表重复一遍——散文一遍、管道表一遍）。"""
    try:
        tables = page.find_tables()
    except Exception:  # noqa: BLE001 — 个别页版面畸形时退回纯文本，不让转换失败
        return [b[4] for b in page.get_text("blocks") if b[4].strip()]
    rects = [t.bbox for t in tables.tables]
    out: list[str] = []
    for b in page.get_text("blocks"):
        if not b[4].strip():
            continue
        x0, y0, x1, y1 = b[:4]
        cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
        inside = any(rx0 <= cx <= rx1 and ry0 <= cy <= ry1 for rx0, ry0, rx1, ry1 in rects)
        if not inside:
            out.append(b[4])
    return out


def convert_pdf(pdf_path: Path, doc_name: str) -> str:
    """文本型 PDF：按页转换，页内表格检测还原为 Markdown 管道表（L1 表格保真）。

    表格不还原时被 ``get_text`` 剥成散文——表头标签与数值的行结构丢失，
    这是英文 10-K 域检索弱的主因之一（FinanceBench 校准定位）。
    图片型（无文本）页面走 D1 OCR 兜底（``ocr.py``，可选依赖 extra=ocr，
    未装则入库一条明确的提示行，绝不静默空文本）。
    """
    from findata.finqa.ocr import ocr_page_text

    doc = pymupdf.open(pdf_path)
    parts = [f"# {doc_name}", ""]
    for i, page in enumerate(doc, start=1):
        table_blocks: list[str] = []
        try:
            tables = page.find_tables()
        except Exception:  # noqa: BLE001 — 同 _page_blocks_outside_tables 的容错口径
            tables = None
        if tables:
            for table in tables.tables:
                rows = [
                    ["" if c is None else str(c) for c in row]
                    for row in (table.extract() or [])
                ]
                md = render_markdown_table(rows)
                if md:
                    table_blocks.append(md)
        text = "\n".join(
            block.strip() for block in _page_blocks_outside_tables(page) if block.strip()
        )
        if not text and not table_blocks and page.get_images():
            text = ocr_page_text(page)  # 扫描页兜底；未装 OCR 返回提示行
        if text or table_blocks:
            parts += ["", f"## Page {i}", ""]
            if text:
                parts.append(text)
            for md in table_blocks:
                parts += ["", md]
    return "\n".join(parts) + "\n"


def convert_docx(docx_path: Path, doc_name: str) -> str:
    import docx  # python-docx

    document = docx.Document(str(docx_path))
    parts = [f"# {doc_name}", ""]
    section = 0
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if (para.style.name or "").startswith("Heading"):
            section += 1
            parts += ["", f"## Heading-{section} {text[:60]}"]
        else:
            parts.append(text)
    return "\n".join(parts) + "\n"


def convert_textlike(text_path: Path, doc_name: str) -> str:
    text = text_path.read_text(encoding="utf-8", errors="replace").strip()
    if "## " not in text:
        # 无分节锚点的纯文本：按 2400 字符插 Section 锚点（供检索层分块）
        parts = [text[i : i + 2400] for i in range(0, len(text), 2400)]
        text = "\n\n".join(
            f"## Section {i}\n\n{part}" for i, part in enumerate(parts, start=1)
        )
    return f"# {doc_name}\n\n{text}\n"


def convert_file(
    path: Path, docs_dir: Path, converted_dir: Path, mask_pii: bool = False
) -> Path:
    """原件已在 docs_dir 下；转换为 converted_dir/<stem>.md 并返回。

    ``mask_pii=True`` 时对**转换后的检索文本**脱敏（身份证/手机号/邮箱/
    银行卡形态数字），原件不动——提问发往模型服务的是脱敏文本，本机
    保留原件原件级隐私（C2）。
    """
    doc_name = safe_stem(path.name)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        md = convert_pdf(path, doc_name)
    elif suffix == ".docx":
        md = convert_docx(path, doc_name)
    elif suffix in (".txt", ".md"):
        md = convert_textlike(path, doc_name)
    else:
        raise ValueError(f"不支持的格式：{suffix}（支持 pdf/docx/txt/md）")
    if mask_pii:
        md, _n = mask_pii_text(md)
    converted_dir.mkdir(parents=True, exist_ok=True)
    out = converted_dir / f"{doc_name}.md"
    out.write_text(md, encoding="utf-8")
    return out


# ── C2 合规：PII 脱敏（上传时可开，默认关）──
# 克制规则：金融文档里 13 位以下的纯数字多半是金额/账号，误伤会毁掉问答
# 价值——只脱敏「身份证 / 手机号 / 邮箱 / 16-19 位连续数字（银行卡长度）
# 与 4×4 分组数字」四类高置信形态。边界如实：增值税号（15/20 位）等
# 长数字串若恰落在 16-19 位会被连带脱敏。
_PII_PATTERNS: list[tuple[str, re.Pattern[str]]] = [
    (
        "身份证",
        re.compile(
            r"\b\d{6}(?:19|20)\d{2}(?:0[1-9]|1[0-2])(?:[0-2]\d|3[01])\d{3}[\dXx]\b"
        ),
    ),
    ("手机号", re.compile(r"\b1[3-9]\d{9}\b")),
    (
        "邮箱",
        re.compile(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b"),
    ),
    ("银行卡(连续)", re.compile(r"\b\d{16,19}\b")),
    ("银行卡(分组)", re.compile(r"\b\d{4} \d{4} \d{4} \d{4}(?: \d{1,3})?\b")),
]


def _mask_one(s: str) -> str:
    """保留前 3 后 2 字符，中间打码（可辨识形态保留，还原不可行）。"""
    if len(s) <= 5:
        return "*" * len(s)
    return s[:3] + "*" * (len(s) - 5) + s[-2:]


def mask_pii_text(text: str) -> tuple[str, int]:
    """对文本做 PII 脱敏，返回（脱敏后文本, 命中数）。确定性、零模型。"""
    total = 0
    for _name, pattern in _PII_PATTERNS:
        text, n = pattern.subn(lambda m: _mask_one(m.group(0)), text)
        total += n
    return text, total


# ── 自由问答 ──

FREE_QA_SYSTEM_PROMPT = (
    "你是严谨的文档问答助手。只能依据提供的文档片段回答，"
    "禁止使用外部知识或常识推测。用与问题相同的语言作答。"
    "你提供的是文档信息整理，不构成投资建议或法律意见。"
)


def _chunk_chars_for(docs: dict[str, str]) -> int:
    """英文主导语料用 2400，否则维持 1200。

    FinanceBench 消融（examples/financebench-recall-v2）：英文财报段落长，
    1200 会切碎金证据段，1200→2400 使 recall@8 翻倍；中文条款语料的
    1200 是 v3 调优值，不动。
    """
    text = "".join(docs.values())
    if not text:
        return MAX_CHUNK_CHARS
    latin = sum(1 for c in text if c.isascii() and c.isalpha())
    cjk = sum(1 for c in text if "\u4e00" <= c <= "\u9fff")
    return 2400 if latin > cjk else MAX_CHUNK_CHARS


def build_free_qa_messages(question: str, chunks_text: str) -> list[dict[str, str]]:
    user = (
        "以下是用户文档中与问题相关的检索片段：\n\n"
        f"{chunks_text}\n\n"
        f"【问题】{question}\n\n"
        "【作答要求】\n"
        "- 只依据片段内容回答，引用出处用【文档编号·Section】标注；\n"
        "- 片段不足以回答时，明确写「提供的文档片段中没有提及」，不许推测；\n"
        "- 回答简洁，不超过 300 字。"
    )
    return [
        {"role": "system", "content": FREE_QA_SYSTEM_PROMPT},
        {"role": "user", "content": user},
    ]


# ── C1 合规：投资决策建议类问题的边界 ──
# 金融场景的产品纪律：文档信息整理 ≠ 投资建议。决策建议型问题（买卖/仓位/
# 价格预测）代码层直接拒答并转引客观事实口径，不依赖模型自觉。
_ADVICE_RE = re.compile(
    r"(该不该|要不要|应不应该|值不值得|能不能|可不可以|可否)"
    r"(买入|卖出|投资|持有|加仓|减仓|清仓|抄底|止盈|止损|买|卖|投)"
    r"|"
    r"(推荐|建议)(我)?(买入|卖出|加仓|减仓|清仓|一只?股票|基金)"
    r"|"
    r"(会涨|会跌|涨还是跌|还能涨|还能跌|看涨|看跌|目标价|能到多少|能赚多少)"
    r"|"
    r"(怎么选股|帮我选股|选哪只|买哪只|哪个更值得投)"
)

ADVICE_REFUSAL = (
    "这个问题在寻求投资决策建议，超出我的服务边界：我只做文档信息整理"
    "（条款、数据、披露事项），不给买卖、仓位或价格预测类建议。"
    "你可以改问文档中的客观事实，例如「文档披露的营收是多少」；"
    "投资决策请咨询持牌投资顾问。"
)


def is_advice_question(question: str) -> bool:
    return bool(_ADVICE_RE.search(question))


# ── C3 合规：引用覆盖率（答案数字被引用片段覆盖的比例） ──
_ANSWER_NUM = re.compile(r"\d[\d,，.]*%?")


def _norm_number(s: str) -> str:
    return s.replace(",", "").replace("，", "")


def answer_numbers(answer: str) -> list[str]:
    """答案中的关键数字（去重保序）。

    噪声控制：滤掉孤立的一位数（「第1条」的 1 这类）；年份、百分比、
    小数与千分位数字保留。归一化剥掉千分位逗号后再比对。
    """
    seen: dict[str, None] = {}
    for m in _ANSWER_NUM.finditer(answer):
        raw = m.group(0)
        norm = _norm_number(raw)
        digits = norm.rstrip("%").replace(".", "")
        if len(digits) < 2 and "%" not in norm and "." not in norm:
            continue
        seen[norm] = None
    return list(seen)


def citation_coverage(
    answer: str, cited_texts: list[str]
) -> tuple[float, list[str]]:
    """答案数字被「已引用片段」文本覆盖的比例（C3 反向归因）。

    M4 归因校验的是「引用是否真实在场」；本函数度量的反向问题——答案里的
    数字是否都能在被引用的片段里找到。覆盖率 < 1 的数字如实列出，
    前端标注「未直接引用原文」，不静默。
    """
    nums = answer_numbers(answer)
    if not nums:
        return 1.0, []
    blob = _norm_number("\n".join(cited_texts))
    uncovered = [n for n in nums if n not in blob]
    return round(1 - len(uncovered) / len(nums), 4), uncovered


# ── 混合检索（L2 深化，产品口径）──
# 默认模型按语料语言二选一（与 _chunk_chars_for 同判据）：fastembed 的
# bge-small 系按语言分模，中英模型都不大（~100MB 量级），本机 CPU 可跑。
EN_EMBED_MODEL = "BAAI/bge-small-en-v1.5"
ZH_EMBED_MODEL = "BAAI/bge-small-zh-v1.5"


def embed_model_for(docs: dict[str, str], override: str = "") -> str:
    if override:
        return override
    text = "".join(docs.values())
    if not text:
        return EN_EMBED_MODEL
    latin = sum(1 for c in text if c.isascii() and c.isalpha())
    cjk = sum(1 for c in text if "\u4e00" <= c <= "\u9fff")
    return EN_EMBED_MODEL if latin > cjk else ZH_EMBED_MODEL


def ask_free(
    question: str,
    docs: dict[str, str],
    client: QwenClient,
    k_total: int = 12,
    embedder: Embedder | None = None,
) -> dict:
    """自由问答：检索 → 单次作答 → 引用归因 + 覆盖率审计 + token 台账。

    C1：决策建议型问题不发模型调用——拒答在代码层，纪律不交给模型自觉。
    L2 深化：``embedder`` 非空时用混合检索（BM25F+向量 RRF，产品口径，
    docs/next-cycle-roadmap.md §3.1.1）；后端不可用**显式回退** BM25F，
    retriever 字段如实标注实际用的臂——回退不是失败，静默才是。
    """
    if is_advice_question(question):
        return {
            "question": question,
            "answer": ADVICE_REFUSAL,
            "citations": [],
            "retrieved_sections": 0,
            "coverage": None,
            "uncovered_numbers": [],
            "refused": True,
            "retriever": "refused",
            "prompt_tokens": 0,
            "completion_tokens": 0,
            "total_tokens": 0,
        }
    chunked = chunk_docs(docs, max_chunk_chars=_chunk_chars_for(docs))
    chunks = None
    retriever = "bm25f"
    if embedder is not None:
        try:
            index: BM25Index | HybridIndex = HybridIndex(chunked, embedder)
            chunks = index.retrieve_many([question], k_per_query=6, total_cap=k_total)
            retriever = f"hybrid:{embedder.name}"
        except EmbeddingUnavailable as exc:
            retriever = f"bm25f(fallback: {exc})"
    if chunks is None:
        chunks = BM25Index(chunked).retrieve_many(
            [question], k_per_query=6, total_cap=k_total
        )
    chunks_text = "\n\n".join(c.block() for c in chunks) or "（文档库中没有命中内容）"
    reply, record = client.chat(
        "free_qa", "free_answer", build_free_qa_messages(question, chunks_text)
    )
    units = [(c.doc_id, c.header, c.text) for c in chunks]
    used = match_citations(extract_citations(reply), [(d, h) for d, h, _ in units])
    answer = reply.strip()
    m = _ANSWER.search(answer)
    if m:  # 模型若按 ANSWER: 格式收尾则取正文
        answer = (answer[: m.start()]).strip() or answer
    cited_texts = [text for i, (_d, _h, text) in enumerate(units) if i in used]
    coverage, uncovered = citation_coverage(answer, cited_texts)
    return {
        "question": question,
        "answer": answer[:2000],
        "citations": [
            {"doc_id": doc_id, "section": header}
            for i, (doc_id, header, _text) in enumerate(units)
            if i in used
        ],
        "retrieved_sections": len(units),
        "coverage": coverage,
        "uncovered_numbers": uncovered,
        "refused": False,
        "retriever": retriever,
        "prompt_tokens": record.prompt_tokens,
        "completion_tokens": record.completion_tokens,
        "total_tokens": record.total_tokens,
    }
